import pandas as pd
import matplotlib.pyplot as plt
from openpyxl.drawing.image import Image

# Read sales data
df = pd.read_csv("sales.csv")

# Calculate total sales
df["Total Sales"] = df["Quantity"] * df["Price"]

# Calculate category-wise sales
category_report = df.groupby("Category")["Total Sales"].sum()

print("Category-wise Sales:")
print(category_report)

# Create chart
category_report.plot(kind="bar")

plt.title("Sales by Category")
plt.xlabel("Category")
plt.ylabel("Total Sales")
plt.tight_layout()

plt.savefig("Sales_by_Category.png")
plt.close()

# Create Excel report
with pd.ExcelWriter("Sales_Report.xlsx") as writer:
    df.to_excel(writer, sheet_name="Raw Data", index=False)
    category_report.to_excel(writer, sheet_name="Category Report")

print("Excel report created successfully!")

# Open the Excel file
from openpyxl import load_workbook

wb = load_workbook("Sales_Report.xlsx")

# Select the Category Report sheet
ws = wb["Category Report"]

# Add the chart image
img = Image("Sales_by_Category.png")
ws.add_image(img, "D2")

# Save the Excel file
wb.save("Sales_Report.xlsx")

print("Chart added to Excel successfully!")

# Add Summary sheet
wb = load_workbook("Sales_Report.xlsx")

ws = wb.create_sheet("Summary")

ws["A1"] = "SALES REPORT SUMMARY"

ws["A3"] = "Total Sales"
ws["B3"] = df["Total Sales"].sum()

ws["A4"] = "Total Quantity"
ws["B4"] = df["Quantity"].sum()

ws["A5"] = "Number of Products"
ws["B5"] = df["Product"].nunique()

ws["A7"] = "Category"
ws["B7"] = "Total Sales"

row = 8

for category, sales in category_report.items():
    ws[f"A{row}"] = category
    ws[f"B{row}"] = sales
    row += 1

wb.save("Sales_Report.xlsx")

print("Summary sheet added successfully!")

# Professional formatting
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = load_workbook("Sales_Report.xlsx")

# Format all sheets
for ws in wb.worksheets:

    # Bold header row
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = max_length + 3

# Format Summary title
ws = wb["Summary"]

ws["A1"].font = Font(bold=True, size=16)
ws["A1"].alignment = Alignment(horizontal="center")

# Format Summary values
for cell in ws["B"]:
    if isinstance(cell.value, (int, float)):
        cell.number_format = '#,##0'

wb.save("Sales_Report.xlsx")

print("Excel formatting completed successfully!")