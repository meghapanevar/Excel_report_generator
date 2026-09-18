import tkinter as tk
from tkinter import filedialog, messagebox
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.drawing.image import Image
import matplotlib.pyplot as plt


def generate_report():

    file_path = filedialog.askopenfilename(
        title="Select Sales CSV File",
        filetypes=[("CSV Files", "*.csv")]
    )

    if not file_path:
        return

    try:

        # -------------------------
        # Read CSV
        # -------------------------

        df = pd.read_csv(file_path)

        # Calculate Total Sales
        df["Total Sales"] = df["Quantity"] * df["Price"]

        # Category-wise Sales
        category_report = (
            df.groupby("Category")["Total Sales"]
            .sum()
            .reset_index()
        )

        # -------------------------
        # Create Chart
        # -------------------------

        plt.figure(figsize=(7, 5))

        plt.bar(
            category_report["Category"],
            category_report["Total Sales"]
        )

        plt.title("Sales by Category")
        plt.xlabel("Category")
        plt.ylabel("Total Sales")

        plt.tight_layout()

        chart_path = "Sales_by_Category.png"

        plt.savefig(chart_path)
        plt.close()

        # -------------------------
        # Create Excel
        # -------------------------

        report_path = "Sales_Report.xlsx"

        with pd.ExcelWriter(report_path) as writer:

            df.to_excel(
                writer,
                sheet_name="Raw Data",
                index=False
            )

            category_report.to_excel(
                writer,
                sheet_name="Category Report",
                index=False
            )

        # Open workbook
        wb = load_workbook(report_path)

        # -------------------------
        # Category Report
        # -------------------------

        ws = wb["Category Report"]

        img = Image(chart_path)

        img.width = 500
        img.height = 300

        ws.add_image(img, "D2")

        # -------------------------
        # Summary Sheet
        # -------------------------

        if "Summary" in wb.sheetnames:
            del wb["Summary"]

        summary = wb.create_sheet("Summary", 0)

        # Dashboard title
        summary.merge_cells("A1:F1")

        summary["A1"] = "SALES ANALYSIS DASHBOARD"

        summary["A1"].font = Font(
            bold=True,
            size=20
        )

        summary["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        summary.row_dimensions[1].height = 35

        # -------------------------
        # KPI Section
        # -------------------------

        summary["A3"] = "TOTAL SALES"
        summary["C3"] = "TOTAL QUANTITY"
        summary["E3"] = "NUMBER OF PRODUCTS"

        summary["A4"] = df["Total Sales"].sum()
        summary["C4"] = df["Quantity"].sum()
        summary["E4"] = df["Product"].nunique()

        for cell in ["A3", "C3", "E3"]:

            summary[cell].font = Font(
                bold=True,
                size=11
            )

            summary[cell].alignment = Alignment(
                horizontal="center"
            )

        for cell in ["A4", "C4", "E4"]:

            summary[cell].font = Font(
                bold=True,
                size=16
            )

            summary[cell].alignment = Alignment(
                horizontal="center"
            )

        # -------------------------
        # Category Section
        # -------------------------

        summary["A7"] = "CATEGORY"
        summary["B7"] = "TOTAL SALES"

        summary["A7"].font = Font(bold=True)
        summary["B7"].font = Font(bold=True)

        row = 8

        for _, data in category_report.iterrows():

            summary[f"A{row}"] = data["Category"]
            summary[f"B{row}"] = data["Total Sales"]

            row += 1

        # -------------------------
        # Add Chart to Summary
        # -------------------------

        summary_chart = Image(chart_path)

        summary_chart.width = 500
        summary_chart.height = 300

        summary.add_image(
            summary_chart,
            "D7"
        )

        # -------------------------
        # Number Formatting
        # -------------------------

        summary["A4"].number_format = "#,##0"
        summary["C4"].number_format = "#,##0"
        summary["E4"].number_format = "#,##0"

        for row in range(
            8,
            8 + len(category_report)
        ):

            summary[f"B{row}"].number_format = "#,##0"

        # -------------------------
        # Column Width
        # -------------------------

        summary.column_dimensions["A"].width = 25
        summary.column_dimensions["B"].width = 18
        summary.column_dimensions["C"].width = 20
        summary.column_dimensions["D"].width = 18
        summary.column_dimensions["E"].width = 22
        summary.column_dimensions["F"].width = 18

        # -------------------------
        # Format All Sheets
        # -------------------------

        for sheet in wb.worksheets:

            for cell in sheet[1]:

                cell.font = Font(
                    bold=True
                )

                cell.alignment = Alignment(
                    horizontal="center"
                )

        # -------------------------
        # Save
        # -------------------------

        wb.save(report_path)

        # Update GUI
        status_label.config(
            text="✓ Report generated successfully!"
        )

        messagebox.showinfo(
            "Success",
            "Excel report generated successfully!\n\n"
            "Sales_Report.xlsx"
        )

    except Exception as e:

        messagebox.showerror(
            "Error",
            str(e)
        )


# =================================
# GUI
# =================================

window = tk.Tk()

window.title("Excel Report Generator")

window.geometry("600x400")

window.resizable(False, False)


# Title

title_label = tk.Label(
    window,
    text="Excel Report Generator",
    font=("Arial", 24, "bold")
)

title_label.pack(pady=40)


# Description

description = tk.Label(
    window,
    text="Upload a CSV file and automatically generate\n"
         "an Excel sales analysis report.",
    font=("Arial", 12)
)

description.pack(pady=10)


# Button

generate_button = tk.Button(
    window,
    text="Upload CSV & Generate Report",
    command=generate_report,
    font=("Arial", 13, "bold"),
    padx=25,
    pady=12
)

generate_button.pack(pady=30)


# Status

status_label = tk.Label(
    window,
    text="Ready to generate report",
    font=("Arial", 11)
)

status_label.pack(pady=10)


# Footer

footer = tk.Label(
    window,
    text="Python • Pandas • OpenPyXL • Matplotlib • Tkinter",
    font=("Arial", 9)
)

footer.pack(side="bottom", pady=20)


# Start GUI

window.mainloop()